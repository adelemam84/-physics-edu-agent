from __future__ import annotations

from math import atan2, cos, pi, sin
from typing import Callable

from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Pt

_SUPPORTED = {
    "force_diagram",
    "graph_plot",
    "particle_model",
    "reaction_profile",
    "cell_structure",
    "food_web",
    "earth_layers",
}


def supported_kinds() -> set[str]:
    return set(_SUPPORTED)


def _xy(rect: tuple, x: float, y: float) -> tuple[int, int]:
    left, top, width, height = rect
    return int(left + width * float(x)), int(top + height * float(y))


def _textbox(slide, x, y, w, h, text: str, *, size: int = 11, bold: bool = False, align=PP_ALIGN.CENTER):
    shape = slide.shapes.add_textbox(int(x), int(y), int(w), int(h))
    frame = shape.text_frame
    frame.clear()
    p = frame.paragraphs[0]
    p.text = str(text or "")
    p.alignment = align
    p.font.size = Pt(size)
    p.font.bold = bold
    return shape


def _line(slide, x1, y1, x2, y2, *, arrow: bool = False):
    shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, int(x1), int(y1), int(x2), int(y2))
    shape.line.width = Pt(1.5)
    if arrow:
        try:
            shape.line.end_arrowhead = True
        except Exception:
            pass
    return shape


def _node(slide, x, y, label: str, *, w=760000, h=360000):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, int(x-w/2), int(y-h/2), int(w), int(h))
    shape.text_frame.clear()
    p = shape.text_frame.paragraphs[0]
    p.text = str(label or "")
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(10)
    return shape


def render_force(slide, params: dict, rect: tuple) -> None:
    ox, oy = _xy(rect, params.get("object_x", .5), params.get("object_y", .5))
    _node(slide, ox, oy, str(params.get("object_label") or ""), w=900000, h=500000)
    for item in params.get("forces") or []:
        x1,y1=_xy(rect,item["x1"],item["y1"]);x2,y2=_xy(rect,item["x2"],item["y2"])
        _line(slide,x1,y1,x2,y2,arrow=True)
        _textbox(slide,(x1+x2)//2-350000,(y1+y2)//2-220000,700000,260000,str(item["label"]),size=9)


def render_graph(slide, params: dict, rect: tuple) -> None:
    left,top,width,height=rect
    x0=left+int(width*.12); y0=top+int(height*.86); xr=left+int(width*.94); yt=top+int(height*.12)
    _line(slide,x0,y0,xr,y0,arrow=True); _line(slide,x0,y0,x0,yt,arrow=True)
    _textbox(slide,xr-650000,y0+50000,650000,260000,str(params["x_label"]),size=9)
    _textbox(slide,x0-700000,yt-100000,650000,260000,str(params["y_label"]),size=9,align=PP_ALIGN.RIGHT)
    xmin,xmax=float(params["x_min"]),float(params["x_max"]); ymin,ymax=float(params["y_min"]),float(params["y_max"])
    def point(p):
        x=x0+int((float(p["x"])-xmin)/(xmax-xmin)*(xr-x0))
        y=y0-int((float(p["y"])-ymin)/(ymax-ymin)*(y0-yt))
        return x,y
    for series in params.get("series") or []:
        pts=[point(p) for p in series.get("points") or []]
        for a,b in zip(pts,pts[1:]): _line(slide,*a,*b)
        for x,y in pts:
            slide.shapes.add_shape(MSO_SHAPE.OVAL,x-65000,y-65000,130000,130000)
        if pts:
            x,y=pts[-1];_textbox(slide,x+70000,y-180000,700000,240000,str(series.get("label") or ""),size=8,align=PP_ALIGN.LEFT)


def render_particle(slide, params: dict, rect: tuple) -> None:
    pos={}
    for p in params.get("particles") or []:
        x,y=_xy(rect,p["x"],p["y"]); pos[str(p["id"])]=(x,y,str(p["label"]))
    for link in params.get("links") or []:
        a=pos.get(str(link["from"])); b=pos.get(str(link["to"]))
        if a and b:_line(slide,a[0],a[1],b[0],b[1])
    for _,(x,y,label) in pos.items():
        shape=slide.shapes.add_shape(MSO_SHAPE.OVAL,x-180000,y-180000,360000,360000)
        shape.text_frame.clear();p=shape.text_frame.paragraphs[0];p.text=label;p.alignment=PP_ALIGN.CENTER;p.font.size=Pt(8)


def render_reaction_profile(slide, params: dict, rect: tuple) -> None:
    left,top,width,height=rect
    x0=left+int(width*.12); y0=top+int(height*.86); xr=left+int(width*.94); yt=top+int(height*.12)
    _line(slide,x0,y0,xr,y0,arrow=True);_line(slide,x0,y0,x0,yt,arrow=True)
    vals=[float(params["reactants_energy"]),float(params["products_energy"]),float(params["transition_energy"])]
    lo=min(vals);hi=max(vals);span=max(1e-9,hi-lo)
    def yy(v):return y0-int((v-lo)/span*(y0-yt)*.72)-int((y0-yt)*.08)
    pts=[(x0+int(width*.10),yy(vals[0])),(x0+int(width*.42),yy(vals[2])),(x0+int(width*.78),yy(vals[1]))]
    _line(slide,*pts[0],*pts[1]);_line(slide,*pts[1],*pts[2])
    _textbox(slide,pts[0][0]-300000,pts[0][1]-260000,900000,240000,str(params["reactants_label"]),size=8)
    _textbox(slide,pts[2][0]-500000,pts[2][1]-260000,1000000,240000,str(params["products_label"]),size=8)
    _textbox(slide,x0-700000,yt-100000,650000,240000,str(params["energy_unit"]),size=8,align=PP_ALIGN.RIGHT)


def render_cell(slide, params: dict, rect: tuple) -> None:
    left,top,width,height=rect
    if params.get("cell_type")=="plant":
        slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,left+int(width*.12),top+int(height*.12),int(width*.68),int(height*.70))
    else:
        slide.shapes.add_shape(MSO_SHAPE.OVAL,left+int(width*.12),top+int(height*.12),int(width*.68),int(height*.70))
    for part in params.get("parts") or []:
        x,y=_xy(rect,.15+float(part["x"])*.62,.16+float(part["y"])*.62)
        slide.shapes.add_shape(MSO_SHAPE.OVAL,x-70000,y-70000,140000,140000)
        _textbox(slide,x+90000,y-130000,900000,260000,str(part["label"]),size=8,align=PP_ALIGN.LEFT)


def render_food_web(slide, params: dict, rect: tuple) -> None:
    pos={}
    for n in params.get("nodes") or []:
        x,y=_xy(rect,n["x"],n["y"]);pos[str(n["id"])]=(x,y,str(n["label"]))
    for edge in params.get("edges") or []:
        a=pos.get(str(edge["from"]));b=pos.get(str(edge["to"]))
        if a and b:
            _line(slide,a[0],a[1],b[0],b[1],arrow=True)
            if edge.get("relation_label"):
                _textbox(slide,(a[0]+b[0])//2-350000,(a[1]+b[1])//2-160000,700000,220000,str(edge["relation_label"]),size=7)
    for _,(x,y,label) in pos.items():_node(slide,x,y,label,w=900000,h=320000)


def render_earth_layers(slide, params: dict, rect: tuple) -> None:
    left,top,width,height=rect
    cx=left+int(width*.38);cy=top+int(height*.52);outer=min(int(width*.27),int(height*.42))
    layers=params.get("layers_outer_to_inner") or []
    total=sum(float(x["relative_thickness"]) for x in layers) or 1.0
    radius=outer
    for i,layer in enumerate(layers):
        slide.shapes.add_shape(MSO_SHAPE.OVAL,cx-radius,cy-radius,radius*2,radius*2)
        _textbox(slide,left+int(width*.72),top+int(height*(.12+i*.10)),int(width*.24),240000,str(layer["label"]),size=8,align=PP_ALIGN.LEFT)
        radius=max(70000,radius-int(outer*float(layer["relative_thickness"])/total))


_RENDERERS: dict[str, Callable] = {
    "force_diagram":render_force,
    "graph_plot":render_graph,
    "particle_model":render_particle,
    "reaction_profile":render_reaction_profile,
    "cell_structure":render_cell,
    "food_web":render_food_web,
    "earth_layers":render_earth_layers,
}


def render_native_science_visual(slide, spec: dict, rect: tuple) -> bool:
    kind=str(spec.get("kind") or "")
    fn=_RENDERERS.get(kind)
    params=spec.get("parameters")
    if not fn or not isinstance(params,dict) or not params:
        return False
    fn(slide,params,rect)
    return True
