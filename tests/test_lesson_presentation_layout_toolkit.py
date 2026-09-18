import unittest
from copy import deepcopy
from io import BytesIO

from pptx import Presentation

from app.lesson_presentation_advanced_design import (
    normalize_design_spec,
    render_presentation_pptx_with_design,
    validate_presentation_advanced_design_edits,
)
from app.services.lesson_presentation_blueprint import build_presentation_blueprint
from app.services.lesson_presentation_editor import presentation_editor_base_hash


class LessonPresentationLayoutToolkitTests(unittest.TestCase):
    REF = "ملف 1: lesson.pdf · صفحة 1"

    def _pack(self):
        return {
            "title": "قانون أوم",
            "subject": "فيزياء",
            "grade_label": "الثالث الثانوي",
            "learning_objectives": ["يطبق قانون أوم"],
            "summary": "يربط القانون بين الجهد والتيار والمقاومة.",
            "sections": [
                {
                    "heading": "الفكرة",
                    "body": "العلاقة V = IR تربط الكميات الكهربائية.",
                    "source_refs": [self.REF],
                }
            ],
            "equations_or_rules": [
                {
                    "label": "قانون أوم",
                    "expression": "V = IR",
                    "notes": "",
                    "source_refs": [self.REF],
                }
            ],
            "worked_examples": [],
            "source_visuals": [],
            "diagram_specs": [
                {
                    "kind": "formula_relationship",
                    "title": "علاقة قانون أوم",
                    "description": "فرق الجهد → شدة التيار → المقاومة",
                    "source_refs": [self.REF],
                }
            ],
            "practice_questions": [
                {
                    "id": "Q1",
                    "prompt": "اختر العلاقة الصحيحة",
                    "answer": "V = IR",
                    "explanation": "مدعومة بالمصدر",
                    "source_refs": [self.REF],
                }
            ],
            "common_mistakes": [],
            "quick_revision": [{"text": "V = IR", "source_refs": [self.REF]}],
        }

    def _blueprint(self):
        return build_presentation_blueprint(
            self._pack(),
            {
                "mode": "lesson_explanation",
                "audience": "teacher",
                "language": "ar",
                "length": "medium",
            },
            source_lesson_pack_id="pack-1",
        )

    def _design(self):
        return {
            "layout": "custom",
            "theme": {
                "background": "#FFFFFF",
                "accent": "#175CD3",
                "title_color": "#101828",
                "text_color": "#344054",
                "font_family": "Arial",
                "content_scale": 1.0,
                "title_align": "right",
                "body_align": "right",
            },
            "elements": {
                "title": {"x": 5, "y": 5, "w": 90, "h": 12},
                "body": {"x": 50, "y": 22, "w": 45, "h": 65},
                "visual": {"x": 5, "y": 22, "w": 40, "h": 65},
            },
            "editor": {
                "grid_size": 5,
                "snap_to_grid": True,
                "show_guides": True,
                "template_id": "compare_split",
                "layers": {
                    "title": {"z": 1, "locked": True},
                    "visual": {"z": 2, "locked": False},
                    "body": {"z": 3, "locked": False},
                },
            },
        }

    def test_toolkit_metadata_is_normalized_and_guarded(self):
        normalized, blockers = normalize_design_spec(self._design())
        self.assertEqual(blockers, [])
        self.assertEqual(normalized["editor"]["grid_size"], 5.0)
        self.assertTrue(normalized["editor"]["snap_to_grid"])
        self.assertTrue(normalized["editor"]["layers"]["title"]["locked"])
        self.assertEqual(normalized["editor"]["layers"]["body"]["z"], 3)

    def test_duplicate_layer_z_values_are_rejected(self):
        spec = self._design()
        spec["editor"]["layers"]["body"]["z"] = 2
        _, blockers = normalize_design_spec(spec)
        self.assertIn("design_layer_order_unique", blockers)

    def test_unknown_template_and_layer_fields_are_rejected(self):
        spec = self._design()
        spec["editor"]["template_id"] = "remote-template"
        spec["editor"]["layers"]["title"]["css"] = "position:fixed"
        _, blockers = normalize_design_spec(spec)
        self.assertIn("design_template_id_invalid", blockers)
        self.assertIn("design_title_layer_invalid", blockers)

    def test_toolkit_edit_requires_reapproval_and_keeps_source_grounding(self):
        base = self._blueprint()
        edited = deepcopy(base)
        edited["slides"][0]["design_spec"] = self._design()
        report = validate_presentation_advanced_design_edits(
            base,
            edited,
            supplied_base_hash=presentation_editor_base_hash(base),
        )
        self.assertTrue(report["ready"], report)
        self.assertTrue(report["teacher_reapproval_required"])
        self.assertTrue(report["resize_handles"])
        self.assertTrue(report["snap_to_grid"])
        self.assertTrue(report["layer_ordering"])
        self.assertTrue(report["element_locking"])
        self.assertTrue(report["copy_paste_style"])
        self.assertTrue(report["reusable_templates"])
        self.assertFalse(report["official_question_bank_write"])
        self.assertTrue(report["content_ingestion_unchanged"])

    def test_pptx_shape_creation_honors_layer_order(self):
        blueprint = self._blueprint()
        target = next(s for s in blueprint["slides"] if s.get("visual_specs"))
        if not target.get("content_blocks"):
            target["content_blocks"] = [{"kind": "summary", "text": "العلاقة V = IR.", "source_refs": [self.REF]}]
        for slide in blueprint["slides"]:
            slide["design_spec"] = self._design()
        data = render_presentation_pptx_with_design(blueprint, "teacher")
        prs = Presentation(BytesIO(data))
        slide_index = blueprint["slides"].index(target)
        slide = prs.slides[slide_index]
        texts = [getattr(shape, "text", "") for shape in slide.shapes]
        title_i = next(i for i, text in enumerate(texts) if target["title"] in text)
        visual_i = next(i for i, text in enumerate(texts) if "فرق الجهد" in text)
        body_text = str(target["content_blocks"][0].get("text") or "")
        body_i = next(i for i, text in enumerate(texts) if body_text and body_text in text)
        self.assertLess(title_i, visual_i)
        self.assertLess(visual_i, body_i)

    def test_workspace_exposes_full_layout_toolkit(self):
        import index
        from app import lesson_presentation_studio_ui

        html = lesson_presentation_studio_ui._workspace("job-1")
        for marker in (
            "Layout Toolkit",
            "toggleResize",
            "toolGrid",
            "toolSnap",
            "toolGuides",
            "layerUp",
            "layerDown",
            "layerLock",
            "copyStyle",
            "pasteStyle",
            "applyTemplate",
            "resize-handle",
        ):
            self.assertIn(marker, html)

    def test_pro_editor_groups_are_validated(self):
        spec = self._design()
        spec["editor"]["groups"] = [["title", "body"]]
        normalized, blockers = normalize_design_spec(spec)
        self.assertEqual(blockers, [])
        self.assertEqual(normalized["editor"]["groups"], [["title", "body"]])

        spec["editor"]["groups"] = [["title", "body"], ["body", "visual"]]
        _, blockers = normalize_design_spec(spec)
        self.assertIn("design_group_overlap", blockers)

    def test_pro_editor_capabilities_require_reapproval(self):
        base = self._blueprint()
        edited = deepcopy(base)
        edited["slides"][0]["design_spec"] = self._design()
        edited["slides"][0]["design_spec"]["editor"]["groups"] = [["title", "body"]]
        report = validate_presentation_advanced_design_edits(
            base,
            edited,
            supplied_base_hash=presentation_editor_base_hash(base),
        )
        self.assertTrue(report["ready"], report)
        for key in (
            "multi_select",
            "group_ungroup",
            "alignment_distribution",
            "keyboard_nudging",
            "zoom_pan",
            "element_inspector",
        ):
            self.assertTrue(report[key])
        self.assertTrue(report["teacher_reapproval_required"])
        self.assertFalse(report["official_question_bank_write"])
        self.assertTrue(report["content_ingestion_unchanged"])

    def test_workspace_exposes_pro_editor_tools(self):
        import index
        from app import lesson_presentation_studio_ui

        html = lesson_presentation_studio_ui._workspace("job-pro")
        for marker in (
            "Pro Editor Tools",
            "proGroup",
            "proUngroup",
            "alignLeft",
            "alignMiddle",
            "distH",
            "distV",
            "proZoom",
            "proPanX",
            "insX",
            "insW",
            "pro-selected",
        ):
            self.assertIn(marker, html)


if __name__ == "__main__":
    unittest.main()
