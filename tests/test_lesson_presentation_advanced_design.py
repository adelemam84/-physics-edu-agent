import unittest
from copy import deepcopy
from io import BytesIO

from pptx import Presentation

from app.lesson_presentation_advanced_design import (
    normalize_design_spec,
    render_presentation_pptx_with_design,
    validate_presentation_advanced_design_edits,
)
from app.lesson_presentation_studio_ui import _workspace
from app.services.lesson_presentation_blueprint import build_presentation_blueprint
from app.services.lesson_presentation_editor import presentation_editor_base_hash
from app.services.lesson_presentation_pptx import pptx_preflight


class LessonPresentationAdvancedDesignTests(unittest.TestCase):
    REF = "ملف 1: lesson.pdf · صفحة 1"

    def _pack(self):
        return {
            "title": "قانون أوم",
            "subject": "فيزياء",
            "grade_label": "الثالث الثانوي",
            "learning_objectives": ["يطبق قانون أوم"],
            "summary": "يربط القانون بين الجهد والتيار والمقاومة.",
            "sections": [{"heading": "الفكرة", "body": "العلاقة V = IR تربط الكميات الكهربائية.", "source_refs": [self.REF]}],
            "equations_or_rules": [{"label": "قانون أوم", "expression": "V = IR", "notes": "", "source_refs": [self.REF]}],
            "worked_examples": [],
            "source_visuals": [],
            "diagram_specs": [{"kind": "formula_relationship", "title": "علاقة قانون أوم", "description": "فرق الجهد → شدة التيار → المقاومة", "source_refs": [self.REF]}],
            "practice_questions": [{"id": "Q1", "prompt": "اختر العلاقة الصحيحة", "answer": "V = IR", "explanation": "مدعومة بالمصدر", "source_refs": [self.REF]}],
            "common_mistakes": [],
            "quick_revision": [{"text": "V = IR", "source_refs": [self.REF]}],
        }

    def _blueprint(self):
        return build_presentation_blueprint(
            self._pack(),
            {"mode": "lesson_explanation", "audience": "teacher", "language": "ar", "length": "medium"},
            source_lesson_pack_id="pack-1",
        )

    def _design(self):
        return {
            "layout": "two_column",
            "theme": {
                "background": "#F8FAFC",
                "accent": "#175CD3",
                "title_color": "#101828",
                "text_color": "#344054",
                "font_family": "Arial",
                "content_scale": 1.1,
                "title_align": "right",
                "body_align": "right",
            },
            "elements": {
                "title": {"x": 6, "y": 5, "w": 88, "h": 13},
                "body": {"x": 52, "y": 23, "w": 42, "h": 62},
                "visual": {"x": 5, "y": 23, "w": 42, "h": 62},
            },
        }

    def test_design_only_edit_is_valid_but_requires_teacher_reapproval(self):
        base = self._blueprint()
        edited = deepcopy(base)
        edited["slides"][0]["design_spec"] = self._design()
        report = validate_presentation_advanced_design_edits(
            base,
            edited,
            supplied_base_hash=presentation_editor_base_hash(base),
        )
        self.assertTrue(report["ready"], report)
        self.assertTrue(report["changed"])
        self.assertTrue(report["teacher_reapproval_required"])
        self.assertTrue(report["advanced_design_editor"])
        self.assertFalse(report["prepared_blueprint"]["approval_state"]["teacher_approved"])
        self.assertEqual(report["prepared_blueprint"]["slides"][0]["design_spec"]["layout"], "two_column")

    def test_out_of_bounds_drag_position_is_rejected(self):
        spec = self._design()
        spec["elements"]["title"] = {"x": 90, "y": 5, "w": 20, "h": 13}
        _, blockers = normalize_design_spec(spec)
        self.assertIn("design_title_rect_bounds", blockers)

    def test_arbitrary_theme_or_css_fields_are_rejected(self):
        spec = self._design()
        spec["theme"]["css"] = "position:fixed"
        _, blockers = normalize_design_spec(spec)
        self.assertIn("design_theme_invalid", blockers)

    def test_source_grounding_stays_immutable_with_design_editor(self):
        base = self._blueprint()
        edited = deepcopy(base)
        target = next(s for s in edited["slides"] if s.get("source_refs"))
        target["design_spec"] = self._design()
        target["source_refs"] = ["forged-source"]
        report = validate_presentation_advanced_design_edits(
            base,
            edited,
            supplied_base_hash=presentation_editor_base_hash(base),
        )
        self.assertFalse(report["ready"])
        self.assertIn("immutable_slide_source_refs", report["blocking_failures"])

    def test_designed_pptx_is_valid_and_editable(self):
        base = self._blueprint()
        for slide in base["slides"]:
            slide["design_spec"] = self._design()
        data = render_presentation_pptx_with_design(base, "teacher")
        report = pptx_preflight(data, len(base["slides"]))
        self.assertTrue(report["ready"], report)
        prs = Presentation(BytesIO(data))
        self.assertEqual(len(prs.slides), len(base["slides"]))
        self.assertGreater(len(prs.slides[0].shapes), 0)

    def test_workspace_exposes_layout_drag_drop_and_theme_controls(self):
        html = _workspace("job-1")
        self.assertIn("Advanced Slide Design", html)
        self.assertIn("designLayout", html)
        self.assertIn("dragDesign", html)
        self.assertIn("designBackground", html)
        self.assertIn("designAccent", html)
        self.assertIn("Noto Sans Arabic", html)
        self.assertIn("pointermove", html)
        self.assertIn("design_spec", html)


if __name__ == "__main__":
    unittest.main()
