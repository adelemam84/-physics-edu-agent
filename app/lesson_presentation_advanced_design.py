from __future__ import annotations

import copy
import json
import re
from io import BytesIO
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from . import lesson_presentation_studio
from . import lesson_presentation_editor_productivity as productivity
from .services import lesson_presentation_editor as base_editor
from .services.lesson_presentation_blueprint import presentation_preflight, project_edition
from .services.lesson_presentation_pptx import _add_notes, pptx_preflight
from .services.presentation_science_visuals import render_native_science_visual

_LAYOUTS = {"standard", "title_left", "two_column", "visual_focus", "split_40_60", "custom"}
_TEMPLATE_IDS = {"clean_standard", "visual_story", "exam_focus", "compare_split", "custom"}
_FONTS = {"Aptos", "Arial", "Noto Sans Arabic"}
_ALIGNS = {"right", "center", "left"}
_HEX = re.compile(r"^#[0-9A-Fa-f]{6}$")
_ROLES = {"title", "body", "visual"}
_GRID_SIZES = {1.0, 2.0, 2.5, 5.0, 10.0}
_DEFAULT_RECTS = {
    "title": {"x": 6.0, "y": 5.0, "w": 88.0, "h": 13.0},
    "body": {"x": 6.0, "y": 21.0, "w": 88.0, "h": 28.0},
    "visual": {"x": 10.0, "y": 52.0, "w": 80.0, "h": 34.0},
}
_DEFAULT_THEME = {
    "background": "#FFFFFF",
    "accent": "#344054",
    "title_color": "#172033",
    "text_color": "#344054",
    "font_family": "Aptos",
    "content_scale": 1.0,
    "title_align": "right",
    "body_align": "right",
}
_DEFAULT_EDITOR = {
    "grid_size": 2.5,
    "snap_to_grid": True,
    "show_guides": True,
    "template_id": "clean_standard",
    "groups": [],
    "layers": {
        "title": {"z": 3, "locked": False},
        "body": {"z": 2, "locked": False},
        "visual": {"z": 1, "locked": False},
    },
}


def _canonical(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))


def _number(value: Any) -> float | None:
    if isinstance(value, bool):
        return None
    try:
        out = float(value)
    except (TypeError, ValueError):
        return None
    return out if out == out else None


def _normalize_rect(role: str, value: Any, blockers: list[str]) -> dict:
    default = copy.deepcopy(_DEFAULT_RECTS[role])
    if value is None:
        return default
    if not isinstance(value, dict) or set(value) - {"x", "y", "w", "h"}:
        blockers.append(f"design_{role}_rect_invalid")
        return default
    rect = {}
    for key in ("x", "y", "w", "h"):
        num = _number(value.get(key))
        if num is None:
            blockers.append(f"design_{role}_rect_invalid")
            return default
        rect[key] = round(num, 3)
    if rect["x"] < 0 or rect["y"] < 0 or rect["w"] < 5 or rect["h"] < 5:
        blockers.append(f"design_{role}_rect_bounds")
    if rect["x"] + rect["w"] > 100 or rect["y"] + rect["h"] > 100:
        blockers.append(f"design_{role}_rect_bounds")
    return rect


def _normalize_editor(value: Any, blockers: list[str]) -> dict:
    if value in (None, {}):
        return copy.deepcopy(_DEFAULT_EDITOR)
    if not isinstance(value, dict) or set(value) - {"grid_size", "snap_to_grid", "show_guides", "template_id", "groups", "layers"}:
        blockers.append("design_editor_invalid")
        return copy.deepcopy(_DEFAULT_EDITOR)

    editor = copy.deepcopy(_DEFAULT_EDITOR)
    if "grid_size" in value:
        grid = _number(value.get("grid_size"))
        if grid not in _GRID_SIZES:
            blockers.append("design_grid_size_invalid")
        else:
            editor["grid_size"] = float(grid)
    for key in ("snap_to_grid", "show_guides"):
        if key in value:
            if not isinstance(value.get(key), bool):
                blockers.append(f"design_{key}_invalid")
            else:
                editor[key] = bool(value[key])
    if "template_id" in value:
        template_id = str(value.get("template_id") or "")
        if template_id not in _TEMPLATE_IDS:
            blockers.append("design_template_id_invalid")
        else:
            editor["template_id"] = template_id

    groups = value.get("groups")
    if groups is not None:
        if not isinstance(groups, list):
            blockers.append("design_groups_invalid")
        else:
            normalized_groups = []
            claimed = set()
            for group in groups:
                if not isinstance(group, list) or len(group) < 2:
                    blockers.append("design_group_invalid")
                    continue
                roles = [str(role) for role in group]
                if any(role not in _ROLES for role in roles) or len(set(roles)) != len(roles):
                    blockers.append("design_group_roles_invalid")
                    continue
                if claimed.intersection(roles):
                    blockers.append("design_group_overlap")
                    continue
                claimed.update(roles)
                normalized_groups.append(roles)
            editor["groups"] = normalized_groups

    layers = value.get("layers")
    if layers is not None:
        if not isinstance(layers, dict) or set(layers) - _ROLES:
            blockers.append("design_layers_invalid")
        else:
            normalized_layers = copy.deepcopy(_DEFAULT_EDITOR["layers"])
            seen_z: list[int] = []
            for role in sorted(_ROLES):
                raw = layers.get(role)
                if raw is None:
                    continue
                if not isinstance(raw, dict) or set(raw) - {"z", "locked"}:
                    blockers.append(f"design_{role}_layer_invalid")
                    continue
                z = raw.get("z", normalized_layers[role]["z"])
                locked = raw.get("locked", normalized_layers[role]["locked"])
                if isinstance(z, bool) or not isinstance(z, int) or not 1 <= z <= 3:
                    blockers.append(f"design_{role}_layer_z_invalid")
                else:
                    normalized_layers[role]["z"] = z
                    seen_z.append(z)
                if not isinstance(locked, bool):
                    blockers.append(f"design_{role}_layer_locked_invalid")
                else:
                    normalized_layers[role]["locked"] = locked
            all_z = [normalized_layers[r]["z"] for r in sorted(_ROLES)]
            if len(set(all_z)) != len(all_z):
                blockers.append("design_layer_order_unique")
            editor["layers"] = normalized_layers
    return editor


def normalize_design_spec(value: Any) -> tuple[dict, list[str]]:
    blockers: list[str] = []
    if value in (None, {}):
        return {
            "layout": "standard",
            "theme": copy.deepcopy(_DEFAULT_THEME),
            "elements": copy.deepcopy(_DEFAULT_RECTS),
            "editor": copy.deepcopy(_DEFAULT_EDITOR),
        }, blockers
    if not isinstance(value, dict):
        return {}, ["design_spec_invalid"]
    if set(value) - {"layout", "theme", "elements", "editor"}:
        blockers.append("design_spec_fields_unsupported")

    layout = str(value.get("layout") or "standard")
    if layout not in _LAYOUTS:
        blockers.append("design_layout_invalid")
        layout = "standard"

    raw_theme = value.get("theme") or {}
    if not isinstance(raw_theme, dict) or set(raw_theme) - set(_DEFAULT_THEME):
        blockers.append("design_theme_invalid")
        raw_theme = {}
    theme = copy.deepcopy(_DEFAULT_THEME)
    for key in ("background", "accent", "title_color", "text_color"):
        if key in raw_theme:
            candidate = str(raw_theme.get(key) or "")
            if not _HEX.fullmatch(candidate):
                blockers.append(f"design_{key}_invalid")
            else:
                theme[key] = candidate.upper()
    if "font_family" in raw_theme:
        font = str(raw_theme.get("font_family") or "")
        if font not in _FONTS:
            blockers.append("design_font_family_invalid")
        else:
            theme["font_family"] = font
    if "content_scale" in raw_theme:
        scale = _number(raw_theme.get("content_scale"))
        if scale is None or not 0.75 <= scale <= 1.35:
            blockers.append("design_content_scale_invalid")
        else:
            theme["content_scale"] = round(scale, 3)
    for key in ("title_align", "body_align"):
        if key in raw_theme:
            align = str(raw_theme.get(key) or "")
            if align not in _ALIGNS:
                blockers.append(f"design_{key}_invalid")
            else:
                theme[key] = align

    raw_elements = value.get("elements") or {}
    if not isinstance(raw_elements, dict) or set(raw_elements) - _ROLES:
        blockers.append("design_elements_invalid")
        raw_elements = {}
    elements = {role: _normalize_rect(role, raw_elements.get(role), blockers) for role in sorted(_ROLES)}
    editor = _normalize_editor(value.get("editor"), blockers)
    return {
        "layout": layout,
        "theme": theme,
        "elements": elements,
        "editor": editor,
    }, list(dict.fromkeys(blockers))


def _without_design(blueprint: dict) -> dict:
    out = copy.deepcopy(blueprint)
    for slide in out.get("slides") or []:
        if isinstance(slide, dict):
            slide.pop("design_spec", None)
    return out


def validate_presentation_advanced_design_edits(
    base_blueprint: dict,
    edited_blueprint: dict,
    *,
    supplied_base_hash: str | None = None,
) -> dict:
    if not isinstance(edited_blueprint, dict):
        return productivity.validate_presentation_productivity_edits(
            base_blueprint,
            edited_blueprint,
            supplied_base_hash=supplied_base_hash,
        )

    blockers: list[str] = []
    normalized_by_id: dict[str, dict] = {}
    for slide in edited_blueprint.get("slides") or []:
        if not isinstance(slide, dict):
            continue
        sid = str(slide.get("slide_id") or "")
        normalized, slide_blockers = normalize_design_spec(slide.get("design_spec"))
        blockers.extend(slide_blockers)
        if sid:
            normalized_by_id[sid] = normalized

    stripped_base = _without_design(base_blueprint)
    stripped_edited = _without_design(edited_blueprint)
    report = productivity.validate_presentation_productivity_edits(
        stripped_base,
        stripped_edited,
        supplied_base_hash=supplied_base_hash,
    )
    blockers.extend(report.get("blocking_failures") or [])

    design_changed = False
    base_by_id = {
        str(x.get("slide_id") or ""): x
        for x in (base_blueprint.get("slides") or [])
        if isinstance(x, dict)
    }
    for slide in edited_blueprint.get("slides") or []:
        if not isinstance(slide, dict):
            continue
        sid = str(slide.get("slide_id") or "")
        before = (base_by_id.get(sid) or {}).get("design_spec")
        if before is None and slide.get("duplicate_of"):
            before = (base_by_id.get(str(slide.get("duplicate_of") or "")) or {}).get("design_spec")
        normalized_before, _ = normalize_design_spec(before)
        if _canonical(normalized_before) != _canonical(normalized_by_id.get(sid)):
            design_changed = True
            break

    prepared = report.get("prepared_blueprint")
    if isinstance(prepared, dict):
        for slide in prepared.get("slides") or []:
            if isinstance(slide, dict):
                sid = str(slide.get("slide_id") or "")
                slide["design_spec"] = copy.deepcopy(normalized_by_id.get(sid) or normalize_design_spec(None)[0])
        if design_changed:
            base_revision = int((base_blueprint.get("approval_state") or {}).get("revision") or 1)
            prepared["approval_state"] = {
                "teacher_approved": False,
                "revision": base_revision + 1,
                "edited": True,
            }

    changed = bool(report.get("changed")) or design_changed
    all_blockers = list(dict.fromkeys(blockers))
    report["ready"] = not all_blockers
    report["blocking_failures"] = all_blockers
    report["changed"] = changed
    report["teacher_reapproval_required"] = changed
    report["prepared_blueprint"] = prepared
    report["edit_digest"] = base_editor.presentation_edit_digest(
        base_editor.presentation_editor_base_hash(base_blueprint),
        edited_blueprint,
    )
    report["advanced_design_editor"] = True
    report["drag_drop_layout"] = True
    report["resize_handles"] = True
    report["snap_to_grid"] = True
    report["layer_ordering"] = True
    report["element_locking"] = True
    report["copy_paste_style"] = True
    report["reusable_templates"] = True
    report["multi_select"] = True
    report["group_ungroup"] = True
    report["alignment_distribution"] = True
    report["keyboard_nudging"] = True
    report["zoom_pan"] = True
    report["element_inspector"] = True
    report["theme_controls"] = True
    report["official_question_bank_write"] = False
    report["content_ingestion_unchanged"] = True
    return report


def _rgb(hex_color: str) -> RGBColor:
    raw = hex_color.lstrip("#")
    return RGBColor(int(raw[0:2], 16), int(raw[2:4], 16), int(raw[4:6], 16))


def _align(value: str) -> PP_ALIGN:
    return {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}.get(
        value,
        PP_ALIGN.RIGHT,
    )


def _rect(spec: dict, role: str) -> tuple[float, float, float, float]:
    r = (spec.get("elements") or {}).get(role) or _DEFAULT_RECTS[role]
    return tuple(float(r[k]) for k in ("x", "y", "w", "h"))


def _inch_rect(spec: dict, role: str) -> tuple:
    x, y, w, h = _rect(spec, role)
    return (
        Inches(13.333333 * x / 100),
        Inches(7.5 * y / 100),
        Inches(13.333333 * w / 100),
        Inches(7.5 * h / 100),
    )


def _text_box(
    slide,
    rect: tuple,
    text: str,
    *,
    size: float,
    color: str,
    font: str,
    align: str,
    bold: bool = False,
):
    shape = slide.shapes.add_textbox(*rect)
    frame = shape.text_frame
    frame.clear()
    p = frame.paragraphs[0]
    p.text = str(text or "")
    p.alignment = _align(align)
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.name = font
    p.font.color.rgb = _rgb(color)
    return shape


def _render_visual(slide, visual: dict, rect: tuple, theme: dict) -> None:
    if render_native_science_visual(slide, visual, rect):
        return
    labels = []
    for key in ("nodes", "steps", "items", "labels"):
        values = visual.get(key)
        if isinstance(values, list):
            for item in values:
                text = (
                    item.get("label") or item.get("title") or item.get("text")
                    if isinstance(item, dict)
                    else item
                )
                if str(text or "").strip():
                    labels.append(str(text).strip())
    if not labels:
        description = str(visual.get("description") or "").strip()
        if description:
            labels = [x.strip() for x in description.replace("→", "|").replace("->", "|").split("|") if x.strip()]
    if not labels:
        labels = [str(visual.get("title") or visual.get("kind") or "رسم توضيحي")]
    labels = labels[:6]
    left, top, width, height = rect
    count = max(1, len(labels))
    gap = width * 0.015
    box_w = (width - gap * (count - 1)) / count
    box_h = min(height * 0.62, Inches(1.2))
    y = top + max(0, (height - box_h) / 2)
    for i, label in enumerate(labels):
        x = left + i * (box_w + gap)
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, box_w, box_h)
        shape.fill.solid()
        shape.fill.fore_color.rgb = _rgb("#FFFFFF")
        shape.line.color.rgb = _rgb(theme["accent"])
        frame = shape.text_frame
        frame.clear()
        p = frame.paragraphs[0]
        p.text = label
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(13 * float(theme["content_scale"]))
        p.font.name = theme["font_family"]
        p.font.color.rgb = _rgb(theme["text_color"])


def _render_table(slide, table_spec: dict, rect: tuple, theme: dict) -> None:
    columns = [str(x) for x in (table_spec.get("columns") or [])]
    rows = [x for x in (table_spec.get("rows") or []) if isinstance(x, dict)][:7]
    if not columns or not rows:
        return
    shape = slide.shapes.add_table(len(rows) + 1, len(columns), *rect)
    table = shape.table
    for c, name in enumerate(columns):
        table.cell(0, c).text = name
    for r, row in enumerate(rows, 1):
        for c, name in enumerate(columns):
            table.cell(r, c).text = str(row.get(name) or "")
    for r, row in enumerate(table.rows):
        for cell in row.cells:
            if r == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = _rgb(theme["accent"])
            for p in cell.text_frame.paragraphs:
                p.alignment = _align(theme["body_align"])
                p.font.size = Pt(11 * float(theme["content_scale"]))
                p.font.name = theme["font_family"]
                p.font.color.rgb = _rgb("#FFFFFF" if r == 0 else theme["text_color"])


def _render_role(slide, role: str, spec: dict, design: dict, theme: dict) -> None:
    scale = float(theme["content_scale"])
    font = theme["font_family"]
    if role == "title":
        _text_box(
            slide,
            _inch_rect(design, "title"),
            spec.get("title") or "",
            size=28 * scale,
            color=theme["title_color"],
            font=font,
            align=theme["title_align"],
            bold=True,
        )
        return
    if role == "body":
        blocks = [x for x in (spec.get("content_blocks") or []) if isinstance(x, dict)]
        if not blocks:
            return
        body = slide.shapes.add_textbox(*_inch_rect(design, "body"))
        frame = body.text_frame
        frame.clear()
        for i, block in enumerate(blocks):
            p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            prefix = "• "
            if block.get("kind") == "answer":
                prefix = "الإجابة: "
            elif block.get("kind") == "explanation":
                prefix = "التفسير: "
            p.text = prefix + str(block.get("text") or "")
            p.alignment = _align(theme["body_align"])
            p.font.size = Pt((17 if len(blocks) <= 5 else 14) * scale)
            p.font.name = font
            p.font.color.rgb = _rgb(theme["text_color"])
        return
    visual_rect = _inch_rect(design, "visual")
    visuals = [x for x in (spec.get("visual_specs") or []) if isinstance(x, dict)]
    tables = [x for x in (spec.get("table_specs") or []) if isinstance(x, dict)]
    if visuals:
        _render_visual(slide, visuals[0], visual_rect, theme)
    elif tables:
        _render_table(slide, tables[0], visual_rect, theme)


def render_presentation_pptx_with_design(blueprint: dict, edition: str) -> bytes:
    has_design = any(
        isinstance(x, dict) and x.get("design_spec")
        for x in (blueprint.get("slides") or [])
    )
    if not has_design:
        return _BASE_RENDER(blueprint, edition)

    projected = project_edition(blueprint, edition)
    report = presentation_preflight(projected, edition=edition)
    if not report.get("ready"):
        raise ValueError(f"presentation preflight failed: {report['blocking_failures']}")

    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)

    for spec in projected.get("slides") or []:
        design, blockers = normalize_design_spec(spec.get("design_spec"))
        if blockers:
            raise ValueError(f"presentation design preflight failed: {blockers}")
        theme = design["theme"]
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        bg = slide.background.fill
        bg.solid()
        bg.fore_color.rgb = _rgb(theme["background"])

        layers = (design.get("editor") or {}).get("layers") or _DEFAULT_EDITOR["layers"]
        ordered_roles = sorted(_ROLES, key=lambda role: int((layers.get(role) or {}).get("z") or 1))
        for role in ordered_roles:
            _render_role(slide, role, spec, design, theme)

        refs = " · ".join(spec.get("source_refs") or [])
        if refs:
            _text_box(
                slide,
                (Inches(.55), Inches(7.02), Inches(12.2), Inches(.26)),
                refs,
                size=7.5,
                color=theme["text_color"],
                font=theme["font_family"],
                align="right",
            )
        if edition == "teacher":
            _add_notes(slide, list(spec.get("speaker_notes") or []))

    stream = BytesIO()
    prs.save(stream)
    return stream.getvalue()


_BASE_RENDER = lesson_presentation_studio.render_presentation_pptx
lesson_presentation_studio.validate_presentation_edits = validate_presentation_advanced_design_edits
lesson_presentation_studio.render_presentation_pptx = render_presentation_pptx_with_design
lesson_presentation_studio.pptx_preflight = pptx_preflight
