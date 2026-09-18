from __future__ import annotations

from io import BytesIO

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from .lesson_presentation_blueprint import presentation_preflight, project_edition
from .presentation_science_visuals import render_native_science_visual


def _set_text(shape, text: str, *, size: int = 24, bold: bool = False) -> None:
    frame = shape.text_frame
    frame.clear()
    p = frame.paragraphs[0]
    p.text = str(text or "")
    p.alignment = PP_ALIGN.RIGHT
    p.font.size = Pt(size)
    p.font.bold = bold


def _add_notes(slide, notes: list[str]) -> None:
    if not notes:
        return
    try:
        notes_slide = slide.notes_slide
        frame = getattr(notes_slide, "notes_text_frame", None)
        if frame is None:
            return
        frame.clear()
        for i, note in enumerate(notes):
            p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            p.text = str(note)
    except Exception:
        # Notes support differs slightly across python-pptx versions; failure to attach
        # notes must not corrupt the editable slide deck.
        return


def _visual_labels(spec: dict) -> list[str]:
    labels = []
    for key in ("nodes", "steps", "items", "labels"):
        values = spec.get(key)
        if isinstance(values, list):
            for value in values:
                if isinstance(value, dict):
                    text = value.get("label") or value.get("title") or value.get("text")
                else:
                    text = value
                if str(text or "").strip():
                    labels.append(str(text).strip())
    if not labels:
        description = str(spec.get("description") or "").strip()
        if description:
            parts = [x.strip() for x in description.replace("→", "|").replace("->", "|").split("|") if x.strip()]
            labels.extend(parts[:5])
    if not labels:
        labels.append(str(spec.get("title") or spec.get("kind") or "رسم توضيحي"))
    return labels[:6]


def _add_visual_diagram(slide, spec: dict, *, top: float = 2.0) -> None:
    """Render a deterministic editable diagram from a source-grounded visual spec."""
    native_rect = (Inches(0.8), Inches(top), Inches(11.7), Inches(2.15))
    if render_native_science_visual(slide, spec, native_rect):
        caption = slide.shapes.add_textbox(Inches(0.9), Inches(top + 2.22), Inches(11.5), Inches(0.34))
        _set_text(caption, f"{spec.get('kind') or 'diagram'} · generated_visual · source-grounded", size=8)
        return
    labels = _visual_labels(spec)
    count = max(1, len(labels))
    left = 0.8
    total_width = 11.7
    gap = 0.18
    box_width = max(1.35, (total_width - gap * (count - 1)) / count)
    for i, label in enumerate(labels):
        x = left + i * (box_width + gap)
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(x),
            Inches(top),
            Inches(box_width),
            Inches(1.25),
        )
        _set_text(shape, label, size=15, bold=(i == 0))
        if i < count - 1:
            arrow_x = x + box_width - 0.02
            arrow = slide.shapes.add_shape(
                MSO_SHAPE.RIGHT_ARROW,
                Inches(arrow_x),
                Inches(top + 0.42),
                Inches(0.42),
                Inches(0.36),
            )
            arrow.text = ""

    kind = str(spec.get("kind") or "diagram")
    caption = slide.shapes.add_textbox(Inches(0.9), Inches(top + 1.48), Inches(11.5), Inches(0.42))
    _set_text(caption, f"{kind} · generated_visual" if spec.get("generated") else kind, size=9)


def _add_table_spec(slide, spec: dict, *, top: float = 1.85) -> None:
    columns = [str(x) for x in (spec.get("columns") or [])]
    rows = [x for x in (spec.get("rows") or []) if isinstance(x, dict)]
    if not columns or not rows:
        return
    rows = rows[:7]
    table_shape = slide.shapes.add_table(
        len(rows) + 1,
        len(columns),
        Inches(0.8),
        Inches(top),
        Inches(11.7),
        Inches(min(4.7, 0.55 * (len(rows) + 1))),
    )
    table = table_shape.table
    for c, name in enumerate(columns):
        table.cell(0, c).text = name
    for r, row in enumerate(rows, 1):
        for c, name in enumerate(columns):
            table.cell(r, c).text = str(row.get(name) or "")
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.text_frame.paragraphs:
                paragraph.alignment = PP_ALIGN.RIGHT
                paragraph.font.size = Pt(12)


def render_presentation_pptx(blueprint: dict, edition: str) -> bytes:
    projected = project_edition(blueprint, edition)
    report = presentation_preflight(projected, edition=edition)
    if not report["ready"]:
        raise ValueError(f"presentation preflight failed: {report['blocking_failures']}")

    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)

    for idx, spec in enumerate(projected.get("slides") or []):
        if idx == 0:
            slide = prs.slides.add_slide(prs.slide_layouts[0])
            _set_text(slide.shapes.title, spec.get("title") or projected.get("title") or "", size=30, bold=True)
            subtitle = slide.placeholders[1]
            _set_text(subtitle, f"{projected.get('subject') or ''} · {projected.get('grade_label') or ''}", size=18)
        else:
            slide = prs.slides.add_slide(prs.slide_layouts[5])
            _set_text(slide.shapes.title, spec.get("title") or "", size=28, bold=True)
            visual_specs = list(spec.get("visual_specs") or [])
            table_specs = list(spec.get("table_specs") or [])
            blocks = list(spec.get("content_blocks") or [])

            if visual_specs:
                _add_visual_diagram(slide, visual_specs[0], top=1.9)
                text_top, text_height = 4.15, 2.25
            elif table_specs:
                _add_table_spec(slide, table_specs[0], top=1.65)
                text_top, text_height = 5.35, 1.05
            else:
                text_top, text_height = 1.45, 4.8

            if blocks:
                body = slide.shapes.add_textbox(Inches(0.8), Inches(text_top), Inches(11.7), Inches(text_height))
                frame = body.text_frame
                frame.clear()
                for i, block in enumerate(blocks):
                    p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
                    prefix = "• "
                    if block.get("kind") == "answer":
                        prefix = "الإجابة: "
                    elif block.get("kind") == "explanation":
                        prefix = "التفسير: "
                    p.text = prefix + str(block.get("text") or "")
                    p.alignment = PP_ALIGN.RIGHT
                    p.font.size = Pt(18 if len(blocks) <= 5 else 15)

            refs = " · ".join(spec.get("source_refs") or [])
            if refs:
                footer = slide.shapes.add_textbox(Inches(0.6), Inches(6.9), Inches(12.0), Inches(0.35))
                _set_text(footer, refs, size=8)
        if edition == "teacher":
            _add_notes(slide, list(spec.get("speaker_notes") or []))

    stream = BytesIO()
    prs.save(stream)
    return stream.getvalue()


def pptx_preflight(data: bytes, expected_slide_count: int) -> dict:
    try:
        prs = Presentation(BytesIO(data))
    except Exception:
        return {"ready": False, "blocking_failures": ["valid_pptx"], "slide_count": 0}
    slide_count = len(prs.slides)
    blockers = []
    if slide_count != int(expected_slide_count):
        blockers.append("slide_count_matches_blueprint")
    if slide_count == 0:
        blockers.append("slides_present")
    return {
        "ready": not blockers,
        "blocking_failures": blockers,
        "slide_count": slide_count,
    }
