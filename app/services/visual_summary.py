from __future__ import annotations

from io import BytesIO
from math import cos, pi, sin
import html
import textwrap

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


def _txt(value: object, limit: int = 220) -> str:
    value = ' '.join(str(value or '').split())
    return value if len(value) <= limit else value[:limit - 1].rstrip() + '…'


def _esc(value: object) -> str:
    return html.escape(str(value or ''))


def _refs(item: dict) -> list[str]:
    return [str(x) for x in (item.get('source_refs') or []) if str(x).strip()]


def build_visual_summary(structured: dict) -> dict:
    sections = []
    for i, item in enumerate(structured.get('sections') or [], 1):
        body = _txt(item.get('body'), 260)
        sections.append({
            'id': f'section-{i}',
            'title': _txt(item.get('heading') or f'قسم {i}', 90),
            'summary': body,
            'source_refs': _refs(item),
            'origin': 'teacher_source',
        })
    equations = [
        {
            'label': _txt(x.get('label'), 80),
            'expression': _txt(x.get('expression'), 120),
            'notes': _txt(x.get('notes'), 160),
            'source_refs': [str(r) for r in (x.get('source_refs') or [])],
            'origin': 'teacher_source',
        }
        for x in (structured.get('equations_or_rules') or [])
    ]
    diagrams = []
    for x in structured.get('diagram_specs') or []:
        provenance = dict(x.get('visual_provenance') or {})
        diagrams.append({
            'title': _txt(x.get('title'), 90),
            'description': _txt(x.get('description'), 180),
            'labels': [_txt(v, 50) for v in (x.get('scientific_labels') or [])],
            'kind': str(x.get('normalized_kind') or x.get('kind') or 'other'),
            'source_refs': [str(r) for r in (x.get('source_refs') or [])],
            'visual_provenance': provenance,
            'review_required': bool((x.get('diagram_engine') or {}).get('review_required')),
        })
    return {
        'title': _txt(structured.get('title') or 'ملخص بصري للدرس', 120),
        'subject': str(structured.get('subject') or ''),
        'grade_label': _txt(structured.get('grade_label'), 80),
        'sections': sections,
        'equations': equations,
        'diagrams': diagrams,
        'summary': _txt(structured.get('summary'), 420),
        'formats': ['mindmap_svg', 'flowchart_svg', 'slides_pptx'],
        'policy': {
            'source_grounded_only': True,
            'preserve_source_refs': True,
            'no_silent_scientific_correction': True,
            'generated_visuals_are_presentation_only': True,
        },
    }


def render_mindmap_svg(summary: dict) -> str:
    sections = list(summary.get('sections') or [])
    width, height = 1200, max(720, 180 + len(sections) * 105)
    cx, cy = width // 2, height // 2
    root_w, root_h = 300, 86
    parts = [
        f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}" direction="rtl">',
        '<rect width="100%" height="100%" fill="#ffffff"/>',
        '<defs><marker id="arrow" markerWidth="10" markerHeight="7" refX="9" refY="3.5" orient="auto"><polygon points="0 0,10 3.5,0 7" fill="#667085"/></marker></defs>',
        f'<rect x="{cx-root_w/2}" y="{cy-root_h/2}" width="{root_w}" height="{root_h}" rx="18" fill="#eff8ff" stroke="#175cd3" stroke-width="2"/>',
        f'<text x="{cx}" y="{cy+7}" text-anchor="middle" font-size="24" font-family="sans-serif" font-weight="700">{_esc(summary.get("title"))}</text>',
    ]
    if not sections:
        sections = [{'title':'لا توجد أقسام منظمة','summary':'','source_refs':[]}]
    left = sections[::2]
    right = sections[1::2]
    for side, items in [(-1, left), (1, right)]:
        for j, item in enumerate(items):
            y = 90 + j * 150
            x = 90 if side < 0 else 760
            w, h = 350, 112
            start_x = cx - root_w/2 if side < 0 else cx + root_w/2
            end_x = x + w if side < 0 else x
            parts.append(f'<line x1="{start_x}" y1="{cy}" x2="{end_x}" y2="{y+h/2}" stroke="#667085" stroke-width="2" marker-end="url(#arrow)"/>')
            parts.append(f'<rect x="{x}" y="{y}" width="{w}" height="{h}" rx="14" fill="#f8fafc" stroke="#d0d5dd"/>')
            parts.append(f'<text x="{x+w/2}" y="{y+32}" text-anchor="middle" font-size="19" font-family="sans-serif" font-weight="700">{_esc(item.get("title"))}</text>')
            body = textwrap.wrap(_txt(item.get('summary'), 110), width=46)[:2]
            for k, line in enumerate(body):
                parts.append(f'<text x="{x+w/2}" y="{y+60+k*21}" text-anchor="middle" font-size="14" font-family="sans-serif">{_esc(line)}</text>')
            refs = '، '.join(item.get('source_refs') or [])
            if refs:
                parts.append(f'<text x="{x+w/2}" y="{y+h-10}" text-anchor="middle" font-size="11" font-family="sans-serif" fill="#667085">{_esc(refs)}</text>')
    parts.append('</svg>')
    return ''.join(parts)


def render_flowchart_svg(summary: dict) -> str:
    sections = list(summary.get('sections') or [])
    width = 1200
    row_h = 150
    height = max(500, 140 + row_h * max(1, len(sections)))
    parts = [
        f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}">',
        '<rect width="100%" height="100%" fill="#ffffff"/>',
        '<defs><marker id="arrow" markerWidth="10" markerHeight="7" refX="9" refY="3.5" orient="auto"><polygon points="0 0,10 3.5,0 7" fill="#475467"/></marker></defs>',
        f'<text x="600" y="48" text-anchor="middle" font-size="28" font-family="sans-serif" font-weight="700">{_esc(summary.get("title"))}</text>',
    ]
    if not sections:
        sections = [{'title':'لا توجد أقسام','summary':'','source_refs':[]}]
    x, w, h = 300, 600, 92
    prev_y = None
    for i, item in enumerate(sections):
        y = 86 + i * row_h
        if prev_y is not None:
            parts.append(f'<line x1="600" y1="{prev_y+h}" x2="600" y2="{y}" stroke="#475467" stroke-width="2.5" marker-end="url(#arrow)"/>')
        parts.append(f'<rect x="{x}" y="{y}" width="{w}" height="{h}" rx="16" fill="#f9fafb" stroke="#98a2b3" stroke-width="1.5"/>')
        parts.append(f'<text x="600" y="{y+30}" text-anchor="middle" font-size="20" font-family="sans-serif" font-weight="700">{_esc(item.get("title"))}</text>')
        parts.append(f'<text x="600" y="{y+57}" text-anchor="middle" font-size="14" font-family="sans-serif">{_esc(_txt(item.get("summary"), 115))}</text>')
        refs = '، '.join(item.get('source_refs') or [])
        if refs:
            parts.append(f'<text x="600" y="{y+78}" text-anchor="middle" font-size="11" font-family="sans-serif" fill="#667085">{_esc(refs)}</text>')
        prev_y = y
    parts.append('</svg>')
    return ''.join(parts)


def _add_title(slide, title: str, subtitle: str = ''):
    box = slide.shapes.add_textbox(Inches(0.7), Inches(0.45), Inches(11.9), Inches(0.75))
    p = box.text_frame.paragraphs[0]
    p.text = title
    p.alignment = PP_ALIGN.RIGHT
    p.font.size = Pt(26)
    p.font.bold = True
    if subtitle:
        sub = slide.shapes.add_textbox(Inches(0.7), Inches(1.15), Inches(11.9), Inches(0.45))
        p2 = sub.text_frame.paragraphs[0]
        p2.text = subtitle
        p2.alignment = PP_ALIGN.RIGHT
        p2.font.size = Pt(12)


def _add_source_footer(slide, refs: list[str]):
    if not refs:
        return
    box = slide.shapes.add_textbox(Inches(0.55), Inches(7.0), Inches(12.2), Inches(0.28))
    p = box.text_frame.paragraphs[0]
    p.text = 'المصدر: ' + '، '.join(refs)
    p.alignment = PP_ALIGN.RIGHT
    p.font.size = Pt(9)


def render_slides_pptx(summary: dict) -> bytes:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    cover = prs.slides.add_slide(blank)
    _add_title(cover, str(summary.get('title') or 'ملخص الدرس'), f"{summary.get('subject','')} · {summary.get('grade_label','')}")
    intro = cover.shapes.add_textbox(Inches(1.0), Inches(2.1), Inches(11.3), Inches(2.5))
    p = intro.text_frame.paragraphs[0]
    p.text = _txt(summary.get('summary'), 420)
    p.alignment = PP_ALIGN.RIGHT
    p.font.size = Pt(22)

    sections = list(summary.get('sections') or [])
    for item in sections:
        slide = prs.slides.add_slide(blank)
        _add_title(slide, item.get('title') or 'فكرة رئيسية')
        body = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.05), Inches(1.7), Inches(11.1), Inches(3.4))
        tf = body.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.text = item.get('summary') or ''
        p.alignment = PP_ALIGN.RIGHT
        p.font.size = Pt(20)
        arrow = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(6.65), Inches(5.15), Inches(6.65), Inches(5.85))
        note = slide.shapes.add_textbox(Inches(3.9), Inches(5.85), Inches(5.5), Inches(0.65))
        np = note.text_frame.paragraphs[0]
        np.text = 'الفكرة التالية مرتبطة بهذا الجزء من تسلسل الدرس'
        np.alignment = PP_ALIGN.CENTER
        np.font.size = Pt(12)
        _add_source_footer(slide, item.get('source_refs') or [])

    equations = list(summary.get('equations') or [])
    if equations:
        slide = prs.slides.add_slide(blank)
        _add_title(slide, 'القوانين والمعادلات')
        y = 1.55
        refs = []
        for eq in equations[:6]:
            box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.25), Inches(y), Inches(10.8), Inches(0.75))
            p = box.text_frame.paragraphs[0]
            p.text = f"{eq.get('label','')}: {eq.get('expression','')}"
            p.alignment = PP_ALIGN.CENTER
            p.font.size = Pt(18)
            refs.extend(eq.get('source_refs') or [])
            y += 0.92
        _add_source_footer(slide, sorted(set(refs)))

    diagrams = list(summary.get('diagrams') or [])
    if diagrams:
        for start in range(0, len(diagrams), 4):
            slide = prs.slides.add_slide(blank)
            _add_title(slide, 'الرسومات والعلاقات البصرية')
            chunk = diagrams[start:start+4]
            refs = []
            for i, d in enumerate(chunk):
                col, row = i % 2, i // 2
                x = 0.75 + col * 6.2
                y = 1.55 + row * 2.55
                box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(5.65), Inches(2.05))
                tf = box.text_frame
                tf.clear()
                p = tf.paragraphs[0]
                p.text = d.get('title') or 'رسم'
                p.font.bold = True
                p.font.size = Pt(17)
                p.alignment = PP_ALIGN.RIGHT
                p2 = tf.add_paragraph()
                p2.text = (d.get('description') or '') + ('\n' + ' ← '.join(d.get('labels') or []) if d.get('labels') else '')
                p2.font.size = Pt(13)
                p2.alignment = PP_ALIGN.RIGHT
                refs.extend(d.get('source_refs') or [])
            _add_source_footer(slide, sorted(set(refs)))

    final = prs.slides.add_slide(blank)
    _add_title(final, 'الخلاصة')
    box = final.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.2), Inches(3.9))
    p = box.text_frame.paragraphs[0]
    p.text = _txt(summary.get('summary'), 520)
    p.alignment = PP_ALIGN.RIGHT
    p.font.size = Pt(22)
    note = final.shapes.add_textbox(Inches(1.0), Inches(6.25), Inches(11.2), Inches(0.45))
    np = note.text_frame.paragraphs[0]
    np.text = 'تم إنشاء هذا العرض من محتوى الدرس المرفوع فقط، دون إضافة علمية صامتة.'
    np.alignment = PP_ALIGN.CENTER
    np.font.size = Pt(10)

    out = BytesIO()
    prs.save(out)
    return out.getvalue()
