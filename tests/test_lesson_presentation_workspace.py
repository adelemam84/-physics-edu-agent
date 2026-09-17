import unittest
from copy import deepcopy
from io import BytesIO

from pptx import Presentation

from app.lesson_presentation_studio_ui import _enhance_pack_preview, _workspace
from app.services.lesson_presentation_blueprint import build_presentation_blueprint
from app.services.lesson_presentation_editor import (
    presentation_editor_base_hash,
    validate_presentation_edits,
)
from app.services.lesson_presentation_pptx import pptx_preflight, render_presentation_pptx


class LessonPresentationWorkspaceTests(unittest.TestCase):
    REF = "ملف 1: lesson.pdf · صفحة 1"

    def _pack(self):
        return {
            "title": "قانون أوم",
            "subject": "فيزياء",
            "grade_label": "الثالث الثانوي",
            "learning_objectives": ["يطبق قانون أوم"],
            "summary": "يربط القانون بين الجهد والتيار والمقاومة.",
            "sections": [
                {"heading": "الفكرة", "body": "العلاقة V = IR تربط الكميات الكهربائية.", "source_refs": [self.REF]}
            ],
            "equations_or_rules": [
                {"label": "قانون أوم", "expression": "V = IR", "notes": "", "source_refs": [self.REF]}
            ],
            "worked_examples": [],
            "source_visuals": [],
            "diagram_specs": [
                {
                    "kind": "formula_relationship",
                    "title": "علاقة قانون أوم",
                    "description": "فرق الجهد → شدة التيار → المقاومة",
                    "source_refs": [self.REF],
                }
            ],
            "practice_questions": [
                {"id": "Q1", "prompt": "اختر العلاقة الصحيحة", "answer": "V = IR", "explanation": "مدعومة بالمصدر", "source_refs": [self.REF]}
            ],
            "common_mistakes": [],
            "quick_revision": [{"text": "V = IR", "source_refs": [self.REF]}],
        }

    def _blueprint(self, audience="student"):
        return build_presentation_blueprint(
            self._pack(),
            {"mode": "lesson_explanation", "audience": audience, "language": "ar", "length": "medium"},
            source_lesson_pack_id="pack-1",
        )

    def test_student_pptx_is_valid_and_matches_blueprint_slide_count(self):
        blueprint = self._blueprint("student")
        data = render_presentation_pptx(blueprint, "student")
        report = pptx_preflight(data, len(blueprint["slides"]))
        self.assertTrue(report["ready"], report)
        self.assertGreater(len(data), 1000)

    def test_teacher_pptx_is_valid(self):
        blueprint = build_presentation_blueprint(
            self._pack(),
            {"mode": "question_driven", "audience": "teacher", "language": "ar", "length": "short"},
            source_lesson_pack_id="pack-1",
        )
        data = render_presentation_pptx(blueprint, "teacher")
        report = pptx_preflight(data, len(blueprint["slides"]))
        self.assertTrue(report["ready"], report)

    def test_diagram_slide_uses_editable_shapes_in_pptx(self):
        blueprint = self._blueprint("student")
        diagram_index = next(i for i, slide in enumerate(blueprint["slides"]) if slide.get("visual_specs"))
        data = render_presentation_pptx(blueprint, "student")
        prs = Presentation(BytesIO(data))
        slide = prs.slides[diagram_index]
        self.assertGreaterEqual(len(slide.shapes), 5)
        all_text = "\n".join(getattr(shape, "text", "") for shape in slide.shapes)
        self.assertIn("فرق الجهد", all_text)
        self.assertIn("generated_visual", all_text)

    def test_editor_allows_reorder_hide_and_text_changes_but_invalidates_approval(self):
        base = self._blueprint("teacher")
        edited = deepcopy(base)
        edited["slides"][0]["title"] = "قانون أوم — تمهيد"
        edited["slides"][1]["hidden"] = True
        edited["slides"][2]["content_blocks"][0]["text"] = "صياغة عرض مختصرة للعلاقة V = IR."
        edited["slides"][0], edited["slides"][-1] = edited["slides"][-1], edited["slides"][0]
        report = validate_presentation_edits(
            base,
            edited,
            supplied_base_hash=presentation_editor_base_hash(base),
        )
        self.assertTrue(report["ready"], report)
        self.assertTrue(report["changed"])
        self.assertTrue(report["teacher_reapproval_required"])
        self.assertFalse(report["prepared_blueprint"]["approval_state"]["teacher_approved"])
        self.assertEqual(report["hidden_slide_count"], 1)
        self.assertEqual(
            report["visible_slide_count"],
            len(base["slides"]) - 1,
        )
        self.assertNotIn("hidden", report["prepared_blueprint"]["slides"][0])

    def test_editor_blocks_source_grounding_mutation(self):
        base = self._blueprint("student")
        edited = deepcopy(base)
        target = next(slide for slide in edited["slides"] if slide.get("source_refs"))
        target["source_refs"] = ["forged-source"]
        report = validate_presentation_edits(
            base,
            edited,
            supplied_base_hash=presentation_editor_base_hash(base),
        )
        self.assertFalse(report["ready"])
        self.assertIn("immutable_slide_source_refs", report["blocking_failures"])

    def test_editor_requires_current_base_hash(self):
        base = self._blueprint("student")
        edited = deepcopy(base)
        edited["slides"][0]["title"] = "عنوان جديد"
        missing = validate_presentation_edits(base, edited)
        stale = validate_presentation_edits(base, edited, supplied_base_hash="stale")
        self.assertIn("editor_base_hash_required", missing["blocking_failures"])
        self.assertIn("stale_editor_base", stale["blocking_failures"])

    def test_hidden_slide_is_excluded_from_renderable_prepared_deck(self):
        base = self._blueprint("student")
        edited = deepcopy(base)
        edited["slides"][-1]["hidden"] = True
        report = validate_presentation_edits(
            base,
            edited,
            supplied_base_hash=presentation_editor_base_hash(base),
        )
        self.assertTrue(report["ready"], report)
        prepared = report["prepared_blueprint"]
        data = render_presentation_pptx(prepared, "student")
        pptx_report = pptx_preflight(data, len(prepared["slides"]))
        self.assertTrue(pptx_report["ready"], pptx_report)
        self.assertEqual(pptx_report["slide_count"], len(base["slides"]) - 1)

    def test_workspace_exposes_interactive_preview_editor_and_guarded_exports(self):
        html = _workspace("job-1")
        self.assertIn("Lesson Presentation Studio", html)
        self.assertIn("presentation/blueprint", html)
        self.assertIn("export-pptx/'+edition", html)
        self.assertIn("20 بطاقة", html)
        self.assertIn("معاينة تفاعلية", html)
        self.assertIn("requestFullscreen", html)
        self.assertIn("ArrowLeft", html)
        self.assertIn("filmstrip", html)
        self.assertIn("Speaker Notes", html)
        self.assertIn("renderDiagram", html)
        self.assertIn("محرر الشرائح داخل المتصفح", html)
        self.assertIn("presentation/editor/preflight", html)
        self.assertIn("validateEdits", html)
        self.assertIn("teacher_reapproved", html)
        self.assertIn("moveUp", html)
        self.assertIn("toggleHidden", html)

    def test_lesson_pack_preview_gets_presentation_link(self):
        source = '<div class=box><a href="/admin/lesson-pack-studio">Lesson Pack Studio</a> · <a href="/admin/dashboard">لوحة التحكم</a></div>'
        html = _enhance_pack_preview(source, "job-1")
        self.assertIn("/admin/lesson-pack-studio/jobs/job-1/presentation", html)

    def test_routes_are_registered(self):
        import index

        paths = {route.path for route in index.app.routes}
        self.assertIn("/admin/lesson-pack-studio/jobs/{job_id}/presentation", paths)
        self.assertIn("/api/admin/lesson-pack-studio/jobs/{job_id}/presentation/editor/preflight", paths)
        self.assertIn("/api/admin/lesson-pack-studio/jobs/{job_id}/presentation/export-pptx/{edition}", paths)


if __name__ == "__main__":
    unittest.main()
