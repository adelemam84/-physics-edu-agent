import inspect
import unittest
from copy import deepcopy
from io import BytesIO

from pptx import Presentation

from app.services.diagram_router import route_diagram
from app.services.lesson_diagram_integrity import diagram_spec_hash
from app.services.lesson_presentation_blueprint import build_presentation_blueprint
from app.services.lesson_presentation_editor import presentation_editor_base_hash, validate_presentation_edits
from app.services.lesson_presentation_pptx import render_presentation_pptx
from app.services.science_diagram_specs import schema_catalog, validate_diagram_spec
from app.services.science_visual_engine_v2 import render_visual_v2


class CrossScienceVisualEngineTests(unittest.TestCase):
    REF = "ملف 1: lesson.pdf · صفحة 1"

    def _pack(self, subject, diagram):
        return {
            "title": "درس علمي",
            "subject": subject,
            "grade_label": "المرحلة الإعدادية",
            "learning_objectives": ["يفسر التمثيل العلمي"],
            "summary": "محتوى مصدر تجريبي.",
            "sections": [{"heading": "الفكرة", "body": "وصف علمي من المصدر.", "source_refs": [self.REF]}],
            "equations_or_rules": [],
            "worked_examples": [],
            "source_visuals": [],
            "diagram_specs": [diagram],
            "practice_questions": [],
            "common_mistakes": [],
            "quick_revision": [{"text": "مراجعة من المصدر.", "source_refs": [self.REF]}],
        }

    def test_router_covers_three_science_domains(self):
        self.assertEqual(route_diagram("force", subject="physics").kind, "force_diagram")
        self.assertEqual(route_diagram("particle", subject="chemistry").kind, "particle_model")
        self.assertEqual(route_diagram("cell", subject="science").kind, "cell_structure")
        self.assertEqual(route_diagram("", "شبكة غذائية", subject="علوم").kind, "food_web")

    def test_schema_catalog_marks_domains(self):
        catalog = schema_catalog()
        self.assertEqual(catalog["kinds"]["force_diagram"]["domain"], "physics")
        self.assertEqual(catalog["kinds"]["particle_model"]["domain"], "chemistry")
        self.assertEqual(catalog["kinds"]["cell_structure"]["domain"], "middle_school_science")
        self.assertEqual(catalog["kinds"]["graph_plot"]["domain"], "cross_science")

    def test_force_diagram_requires_explicit_vectors(self):
        valid = validate_diagram_spec("force_diagram", {
            "object_label": "جسم",
            "object_x": .5,
            "object_y": .5,
            "forces": [{"label": "F", "x1": .5, "y1": .5, "x2": .8, "y2": .5}],
        })
        self.assertTrue(valid["valid"], valid)
        invalid = validate_diagram_spec("force_diagram", {"object_label": "جسم", "forces": []})
        self.assertFalse(invalid["valid"])

    def test_chemistry_reaction_profile_rejects_impossible_transition_input(self):
        report = validate_diagram_spec("reaction_profile", {
            "reactants_label": "R",
            "products_label": "P",
            "energy_unit": "kJ/mol",
            "reactants_energy": 20,
            "products_energy": 10,
            "transition_energy": 5,
        })
        self.assertFalse(report["valid"])

    def test_middle_school_food_web_requires_known_nodes(self):
        report = validate_diagram_spec("food_web", {
            "nodes": [
                {"id": "a", "label": "نبات", "x": .2, "y": .7},
                {"id": "b", "label": "حشرة", "x": .5, "y": .4},
            ],
            "edges": [{"from": "a", "to": "missing"}],
        })
        self.assertFalse(report["valid"])

    def test_renderers_are_deterministic_and_review_gated(self):
        params = {
            "particles": [
                {"id": "a", "label": "A", "x": .3, "y": .5},
                {"id": "b", "label": "B", "x": .7, "y": .5},
            ],
            "links": [{"from": "a", "to": "b"}],
        }
        normalized = validate_diagram_spec("particle_model", params)
        self.assertTrue(normalized["valid"])
        a = render_visual_v2("particle_model", "نموذج", normalized["normalized"])
        b = render_visual_v2("particle_model", "نموذج", normalized["normalized"])
        self.assertEqual(a["svg"], b["svg"])
        self.assertTrue(a["review_required"])
        self.assertTrue(a["deterministic"])

    def test_presentation_blueprint_carries_visual_parameters_immutably(self):
        diagram = {
            "kind": "cell_structure",
            "title": "تركيب الخلية",
            "description": "الأجزاء من المصدر",
            "source_refs": [self.REF],
            "scientific_labels": ["النواة"],
            "parameters": {"cell_type": "animal", "parts": [{"label": "النواة", "x": .5, "y": .5}]},
            "routing": {"kind": "cell_structure"},
            "diagram_engine": {"valid": True, "svg": "<svg></svg>", "domain": "middle_school_science"},
        }
        pack = self._pack("علوم", diagram)
        bp = build_presentation_blueprint(
            pack,
            {"mode": "lesson_explanation", "audience": "teacher", "language": "ar", "length": "medium"},
            source_lesson_pack_id="pack-science",
        )
        slide = next(s for s in bp["slides"] if s.get("visual_specs"))
        visual = slide["visual_specs"][0]
        self.assertEqual(visual["parameters"]["cell_type"], "animal")
        edited = deepcopy(bp)
        target = next(s for s in edited["slides"] if s.get("visual_specs"))
        target["visual_specs"][0]["parameters"]["cell_type"] = "plant"
        report = validate_presentation_edits(
            bp,
            edited,
            supplied_base_hash=presentation_editor_base_hash(bp),
        )
        self.assertFalse(report["ready"])
        self.assertIn("immutable_visual_parameters", report["blocking_failures"])

    def test_pptx_uses_editable_native_shapes_for_cross_science_visual(self):
        params = {
            "object_label": "جسم",
            "object_x": .5,
            "object_y": .5,
            "forces": [
                {"label": "F", "x1": .5, "y1": .5, "x2": .8, "y2": .5},
                {"label": "W", "x1": .5, "y1": .5, "x2": .5, "y2": .85},
            ],
        }
        diagram = {
            "kind": "force_diagram",
            "title": "القوى",
            "description": "قوى من المصدر",
            "source_refs": [self.REF],
            "parameters": params,
            "diagram_engine": {
                "schema_validated": True,
                "teacher_reviewed": True,
                "review_required": False,
                "approved_spec_hash": diagram_spec_hash({"kind": "force_diagram", "parameters": params}),
            },
        }
        bp = build_presentation_blueprint(
            self._pack("فيزياء", diagram),
            {"mode": "lesson_explanation", "audience": "teacher", "language": "ar", "length": "medium"},
            source_lesson_pack_id="pack-physics",
        )
        data = render_presentation_pptx(bp, "teacher")
        prs = Presentation(BytesIO(data))
        visual_slide = next(s for s in prs.slides if any(getattr(sh, "text", "") == "جسم" for sh in s.shapes))
        texts = [getattr(sh, "text", "") for sh in visual_slide.shapes]
        self.assertIn("جسم", texts)
        self.assertIn("F", texts)
        self.assertIn("W", texts)

    def test_editor_declares_science_visual_metadata_immutable(self):
        from app.services import lesson_presentation_editor
        src = inspect.getsource(lesson_presentation_editor)
        for field in ("scientific_labels", "parameters", "routing", "diagram_engine"):
            self.assertIn(f'"{field}"', src)


if __name__ == "__main__":
    unittest.main()
